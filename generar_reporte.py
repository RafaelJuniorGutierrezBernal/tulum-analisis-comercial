"""Genera el reporte comercial de Tulum a partir de los CSV de la carpeta data/.

Uso (desde cualquier carpeta):
    python generar_reporte.py

Entrada:  data/ventas.csv, data/productos.csv, data/clientes.csv, data/cotizaciones.csv
Salida:   data/reporte_tulum.xlsx
"""
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

BASE = Path(__file__).resolve().parent
DATA = BASE / "data"
SALIDA = DATA / "reporte_tulum.xlsx"

FMT_NUMERO = "#,##0"
FMT_PORCENTAJE = '0.0"%"'  # los porcentajes se guardan en escala 0-100


def cargar():
    v = pd.read_csv(DATA / "ventas.csv", parse_dates=["fecha"])
    p = pd.read_csv(DATA / "productos.csv")
    c = pd.read_csv(DATA / "clientes.csv")
    q = pd.read_csv(DATA / "cotizaciones.csv", parse_dates=["fecha"])
    for df in (v, p, c, q):
        df.columns = df.columns.str.strip().str.lower()
    # "nombre" existe en productos y en clientes: se renombra para evitar colisiones al unir
    c = c.rename(columns={"nombre": "nombre_cliente"})
    return v, p, c, q


def unir(base, otra, clave, nombre_tabla):
    """Left join que falla si alguna fila de base no encuentra match en otra."""
    res = base.merge(otra, on=clave, how="left", validate="many_to_one")
    columnas_nuevas = [col for col in otra.columns if col != clave]
    sin_match = res[columnas_nuevas].isna().any(axis=1).sum()
    if sin_match:
        raise ValueError(f"{sin_match} filas de ventas sin match en {nombre_tabla} (clave: {clave})")
    return res


def controles(v, p, c, q):
    filas = [(n, len(df), int(df.isna().sum().sum()), int(df.duplicated().sum()))
             for n, df in [("ventas", v), ("productos", p), ("clientes", c), ("cotizaciones", q)]]
    return pd.DataFrame(filas, columns=["tabla", "filas", "nulos", "duplicados"])


def analizar(v, p, c, q):
    d = unir(v, p, "referencia", "productos")
    d = unir(d, c, "cliente_id", "clientes")

    d["ingreso"] = d.cantidad * d.precio_venta_unitario
    d["costo"] = d.cantidad * d.costo_unitario
    d["margen"] = d.ingreso - d.costo
    d["retraso_dias"] = d.dias_produccion_reales - d.dias_produccion_prometidos
    d["mes"] = d.fecha.dt.to_period("M").astype(str)

    por_ref = (d.groupby(["referencia", "nombre", "categoria"])
               .agg(unidades=("cantidad", "sum"), ingreso=("ingreso", "sum"), margen=("margen", "sum"),
                    devueltas=("unidades_devueltas", "sum"))
               .assign(margen_pct=lambda x: (x.margen / x.ingreso * 100).round(2),
                       tasa_devolucion_pct=lambda x: (x.devueltas / x.unidades * 100).round(2))
               .sort_values("margen", ascending=False).reset_index())

    por_cat = (d.groupby("categoria").agg(ingreso=("ingreso", "sum"), margen=("margen", "sum"))
               .assign(margen_pct=lambda x: (x.margen / x.ingreso * 100).round(2)).reset_index())

    mensual = d.groupby("mes").agg(ingreso=("ingreso", "sum"), margen=("margen", "sum")).reset_index()
    mensual["variacion_pct"] = (mensual.ingreso.pct_change() * 100).round(1)

    cli = (d.groupby(["cliente_id", "ciudad", "canal_contacto"])
           .agg(pedidos=("pedido_id", "nunique"), ingreso=("ingreso", "sum"), margen=("margen", "sum"))
           .sort_values("ingreso", ascending=False).reset_index())
    cli["pct_acumulado"] = (cli.ingreso.cumsum() / cli.ingreso.sum() * 100).round(1)
    cli["recompra"] = cli.pedidos > 1

    tiempos = (d.groupby("categoria").agg(prometido=("dias_produccion_prometidos", "mean"),
               real=("dias_produccion_reales", "mean"), retraso_prom=("retraso_dias", "mean"),
               pct_con_retraso=("retraso_dias", lambda s: (s > 0).mean() * 100)).round(1).reset_index())

    estados = q.groupby("estado").size().rename("cotizaciones").reset_index()
    perdidas = (q[q.estado == "Perdida"].groupby("motivo_perdida").size()
                .rename("cotizaciones").sort_values(ascending=False).reset_index())
    tasa_conversion = (q.estado == "Convertida").mean() * 100

    canal = (d.groupby("canal_contacto").agg(ingreso=("ingreso", "sum"), pedidos=("pedido_id", "nunique"))
             .reset_index().sort_values("ingreso", ascending=False))

    ingreso_total = d.ingreso.sum()
    margen_total = d.margen.sum()
    resumen = pd.DataFrame({"indicador": [
        "Ingreso total", "Margen total", "Margen %", "Pedidos", "Clientes activos",
        "Clientes que recompran %", "Top 20% clientes aportan % del ingreso",
        "Pedidos con retraso %", "Conversion de cotizaciones %"],
        "valor": [ingreso_total, margen_total, round(margen_total / ingreso_total * 100, 2),
                  d.pedido_id.nunique(), cli.shape[0], round(cli.recompra.mean() * 100, 1),
                  round(cli.head(max(1, int(len(cli) * .2))).ingreso.sum() / cli.ingreso.sum() * 100, 1),
                  round((d.retraso_dias > 0).mean() * 100, 1), round(tasa_conversion, 1)]})

    return dict(Resumen=resumen, Por_referencia=por_ref, Por_categoria=por_cat, Mensual=mensual,
                Clientes=cli, Tiempos=tiempos, Cotizaciones_estado=estados,
                Cotizaciones_perdidas=perdidas, Canales=canal)


def _formato_resumen(ws):
    """En la hoja Resumen la columna 'valor' mezcla montos, conteos y porcentajes: se formatea fila por fila."""
    for fila in ws.iter_rows(min_row=2):
        indicador, valor = fila[0].value, fila[1]
        valor.number_format = FMT_PORCENTAJE if "%" in str(indicador) else FMT_NUMERO


def formatear(ruta):
    wb = load_workbook(ruta)
    for ws in wb:
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")

        for col in ws.columns:
            letra = get_column_letter(col[0].column)
            ws.column_dimensions[letra].width = max(len(str(x.value or "")) for x in col) + 3
            nombre = str(col[0].value)
            for x in col[1:]:
                if "pct" in nombre:
                    x.number_format = FMT_PORCENTAJE
                elif nombre in ("ingreso", "margen", "costo"):
                    x.number_format = FMT_NUMERO

        if ws.title == "Resumen":
            _formato_resumen(ws)
        ws.freeze_panes = "A2"
    wb.save(ruta)


def main():
    v, p, c, q = cargar()
    hojas = analizar(v, p, c, q)
    hojas["Controles"] = controles(v, p, c, q)
    SALIDA.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(SALIDA, engine="openpyxl") as w:
        for nombre, df in hojas.items():
            df.to_excel(w, sheet_name=nombre[:31], index=False)
    formatear(SALIDA)
    print(f"Reporte generado: {SALIDA}")


if __name__ == "__main__":
    main()
